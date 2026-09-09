"""Offset ranges must be checked before datetime can normalize invalid input."""
import json
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from demo_fixtures import CLOCK, prepare_demo
from input_contract import ContractError, timestamp

ROOT = Path(__file__).resolve().parent


@pytest.mark.parametrize('offset', [
    '+00:60', '+00:99', '-00:60', '-00:99',
    '+23:60', '-23:60', '+24:00', '-24:00',
])
def test_out_of_range_offset_rejected(offset):
    with pytest.raises(ContractError, match='INVALID_TIME'):
        timestamp('2030-01-01T12:59:00' + offset)


@pytest.mark.parametrize('value, expected', [
    ('2030-01-01T12:59:00+00:59', '2030-01-01T12:00:00+00:00'),
    ('2030-01-01T12:59:00-00:59', '2030-01-01T13:58:00+00:00'),
    ('2030-01-01T12:59:00+23:59', '2029-12-31T13:00:00+00:00'),
    ('2030-01-01T12:59:00-23:59', '2030-01-02T12:58:00+00:00'),
    ('2030-01-01T12:59:00Z', '2030-01-01T12:59:00+00:00'),
])
def test_legal_offset_preserves_conversion(value, expected):
    assert timestamp(value).isoformat() == expected


@pytest.mark.parametrize('source', ['excel', 'sqlite'])
@pytest.mark.parametrize('field', ['observed', 'as_of'])
@pytest.mark.parametrize('offset', ['+00:60', '-00:99'])
def test_cli_rejects_invalid_offset_without_optional_checker(tmp_path, source, field, offset):
    fixtures = tmp_path / 'synthetic inputs'
    prepare_demo(fixtures)
    bad = '2030-01-01T12:59:00' + offset
    if field == 'observed':
        if source == 'sqlite':
            conn = sqlite3.connect(fixtures / 'input.sqlite3')
            try:
                conn.execute('UPDATE observations SET sampled=?', (bad,))
                conn.commit()
            finally:
                conn.close()
        else:
            book = load_workbook(fixtures / 'input.xlsx')
            try:
                for row in book['Observations'].iter_rows(min_row=2):
                    row[3].value = bad
                book.save(fixtures / 'input.xlsx')
            finally:
                book.close()
    input_file = fixtures / ('input.xlsx' if source == 'excel' else 'input.sqlite3')
    before = input_file.read_bytes()
    output = tmp_path / 'result'
    # Emulate an installation without optional date-time validation in a new
    # process, then execute the real CLI entry with its full explicit arguments.
    launcher = (
        "import sys,runpy; from jsonschema import FormatChecker; "
        "FormatChecker.checkers.pop('date-time', None); "
        "sys.path.insert(0, sys.argv[1]); "
        "p=sys.argv.pop(1)+'/cli.py'; runpy.run_path(p,run_name='__main__')"
    )
    result = subprocess.run([
        sys.executable, '-B', '-c', launcher, str(ROOT),
        '--source', source, '--input', str(input_file),
        '--model', str(fixtures / 'model.json'),
        '--objects', str(fixtures / 'objects.json'),
        '--binding', str(fixtures / 'business.json'),
        '--mapping', str(fixtures / (source + '-mapping.json')),
        '--collection', 'Observations' if source == 'excel' else 'observations',
        '--as-of', bad if field == 'as_of' else CLOCK,
        '--output', str(output),
    ], capture_output=True, text=True, encoding='utf-8', timeout=30,
        env=dict(os.environ, PYTHONUTF8='1'))
    assert result.returncode == 2, result.stdout + result.stderr
    assert json.loads(result.stdout)['code'] == 'INVALID_TIME'
    assert not output.exists()
    assert input_file.read_bytes() == before
