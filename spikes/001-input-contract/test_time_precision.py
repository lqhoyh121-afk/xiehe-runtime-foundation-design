"""Never silently truncate input timestamps beyond the supported precision."""
import json
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from demo_fixtures import CLOCK, prepare_demo
from input_contract import ContractError, SQLiteSource, load_snapshot, read_json, timestamp

ROOT = Path(__file__).resolve().parent


@pytest.mark.parametrize('fraction', ['0000001', '1234567', '0000000', '123456789'])
@pytest.mark.parametrize('zone', ['Z', '+08:00', '-00:30'])
def test_excess_precision_is_not_silently_truncated(fraction, zone):
    with pytest.raises(ContractError, match='INVALID_TIME'):
        timestamp('2030-01-01T12:00:00.' + fraction + zone)


@pytest.mark.parametrize('value, expected', [
    ('2030-01-01T12:00:00Z', '2030-01-01T12:00:00+00:00'),
    ('2030-01-01T12:00:00.1Z', '2030-01-01T12:00:00.100000+00:00'),
    ('2030-01-01T12:00:00.000001Z', '2030-01-01T12:00:00.000001+00:00'),
    ('2030-01-01T20:00:00.123456+08:00', '2030-01-01T12:00:00.123456+00:00'),
    ('2030-01-01T11:30:00.999999-00:30', '2030-01-01T12:00:00.999999+00:00'),
])
def test_supported_precision_is_preserved(value, expected):
    assert timestamp(value).isoformat() == expected


@pytest.mark.parametrize('observed, code', [
    ('2030-01-01T12:00:00.000001Z', 'FUTURE_OBSERVATION'),
    ('2030-01-01T11:54:59.999999Z', 'STALE_FACT'),
    ('2030-01-01T12:00:00.000000Z', None),
    ('2030-01-01T11:55:00.000000Z', None),
])
def test_microsecond_freshness_boundaries(tmp_path, observed, code):
    fixtures = tmp_path / 'freshness'
    prepare_demo(fixtures)
    conn = sqlite3.connect(fixtures / 'input.sqlite3')
    try:
        conn.execute('UPDATE observations SET sampled=?', (observed,))
        conn.commit()
    finally:
        conn.close()
    source = SQLiteSource(fixtures / 'input.sqlite3', 'observations',
                          read_json(fixtures / 'sqlite-mapping.json'))
    args = (fixtures / 'model.json', fixtures / 'objects.json', source,
            read_json(fixtures / 'business.json')['model_ref'], timestamp(CLOCK))
    if code:
        with pytest.raises(ContractError, match=code):
            load_snapshot(*args)
    else:
        snapshot = load_snapshot(*args)
        assert len(snapshot.to_dict()['records']) == 2


@pytest.mark.parametrize('source', ['excel', 'sqlite'])
@pytest.mark.parametrize('field', ['observed', 'as_of'])
def test_cli_rejects_precision_loss_without_optional_checker(tmp_path, source, field):
    fixtures = tmp_path / '合成时间 inputs'
    prepare_demo(fixtures)
    bad = '2030-01-01T12:00:00.0000001Z'
    if field == 'observed':
        if source == 'sqlite':
            conn = sqlite3.connect(fixtures / 'input.sqlite3')
            try:
                conn.execute('UPDATE observations SET sampled=?', (bad,))
                conn.commit()
            finally:
                conn.close()
        else:
            book = load_workbook(fixtures / 'input.xlsx')
            try:
                for row in book['Observations'].iter_rows(min_row=2):
                    row[3].value = bad
                book.save(fixtures / 'input.xlsx')
            finally:
                book.close()
    before = {p.name: p.read_bytes() for p in fixtures.iterdir()}
    output = tmp_path / 'result'
    launcher = (
        "import sys,runpy; from jsonschema import FormatChecker; "
        "FormatChecker.checkers.pop('date-time', None); "
        "sys.path.insert(0, sys.argv[1]); "
        "p=sys.argv.pop(1)+'/cli.py'; runpy.run_path(p,run_name='__main__')"
    )
    result = subprocess.run([
        sys.executable, '-B', '-c', launcher, str(ROOT),
        '--source', source,
        '--input', str(fixtures / ('input.xlsx' if source == 'excel' else 'input.sqlite3')),
        '--model', str(fixtures / 'model.json'), '--objects', str(fixtures / 'objects.json'),
        '--binding', str(fixtures / 'business.json'),
        '--mapping', str(fixtures / (source + '-mapping.json')),
        '--collection', 'Observations' if source == 'excel' else 'observations',
        '--as-of', bad if field == 'as_of' else CLOCK, '--output', str(output),
    ], capture_output=True, text=True, encoding='utf-8', timeout=30,
        env=dict(os.environ, PYTHONUTF8='1'))
    assert result.returncode == 2, result.stdout + result.stderr
    assert json.loads(result.stdout)['code'] == 'INVALID_TIME'
    assert 'Traceback' not in result.stderr
    assert not output.exists()
    assert {p.name: p.read_bytes() for p in fixtures.iterdir()} == before
