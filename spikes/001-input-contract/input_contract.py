"""SPIKE ONLY: one observation model, interchangeable local read adapters."""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook

MAX_ROWS = 1000


class ContractError(ValueError):
    def __init__(self, code, detail=''):
        self.code = code
        super().__init__(f'{code}: {detail}')


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'), allow_nan=False)


def digest(value):
    return hashlib.sha256(canonical(value).encode('utf-8')).hexdigest()


def read_json(path):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ContractError('DUPLICATE_JSON_KEY', key)
            result[key] = value
        return result
    try:
        return json.loads(Path(path).read_text(encoding='utf-8-sig'), object_pairs_hook=pairs,
                          parse_constant=lambda x: (_ for _ in ()).throw(ContractError('NONFINITE_NUMBER')))
    except (OSError, json.JSONDecodeError) as exc:
        raise ContractError('JSON_READ_FAILED', type(exc).__name__) from exc


def timestamp(value):
    if not isinstance(value, str) or not re.fullmatch(
        r'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})?', value
    ):
        raise ContractError('INVALID_TIME')
    try:
        result = datetime.fromisoformat(value.replace('Z', '+00:00'))
    except (ValueError, AttributeError, TypeError) as exc:
        raise ContractError('INVALID_TIME') from exc
    if result.tzinfo is None or result.utcoffset() is None:
        raise ContractError('TIMEZONE_REQUIRED')
    try:
        return result.astimezone(timezone.utc)
    except OverflowError as exc:
        raise ContractError('INVALID_TIME') from exc


@dataclass(frozen=True)
class SourceBatch:
    rows: list
    provenance: dict


@dataclass(frozen=True)
class ExcelSource:
    path: Path
    sheet: str
    mapping: dict

    def read(self):
        book = None
        try:
            raw = Path(self.path).read_bytes()
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            sheet = book[self.sheet]
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            headers = [c.value for c in first]
            if len(headers) != len(set(headers)):
                raise ContractError('DUPLICATE_HEADER')
            if any(name not in headers for name in self.mapping.values()):
                raise ContractError('MISSING_COLUMN')
            indexes = {key: headers.index(name) for key, name in self.mapping.items()}
            rows = []
            for row in iterator:
                cells = {key: row[index] for key, index in indexes.items()}
                if all(c.value is None for c in cells.values()):
                    continue
                if any(c.data_type == 'f' for c in cells.values()):
                    raise ContractError('FORMULA_NOT_ALLOWED')
                rows.append({key: c.value for key, c in cells.items()})
                if len(rows) > MAX_ROWS:
                    raise ContractError('BATCH_TOO_LARGE')
            return SourceBatch(rows, {'kind': 'excel', 'locator': str(Path(self.path).resolve()),
                                      'sheet': self.sheet, 'mapping': self.mapping,
                                      'sha256': hashlib.sha256(raw).hexdigest(), 'digest_scope': 'file_bytes'})
        except ContractError:
            raise
        except Exception as exc:
            raise ContractError('SOURCE_READ_FAILED', type(exc).__name__) from exc
        finally:
            if book is not None:
                book.close()


@dataclass(frozen=True)
class SQLiteSource:
    path: Path
    table: str
    mapping: dict

    def read(self):
        names = [self.table, *self.mapping.values()]
        if not all(isinstance(n, str) and re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', n) for n in names):
            raise ContractError('INVALID_SQL_IDENTIFIER')
        conn = None
        try:
            conn = sqlite3.connect(Path(self.path).resolve().as_uri() + '?mode=ro', uri=True)
            conn.execute('PRAGMA query_only=ON')
            columns = ','.join('"' + n + '"' for n in self.mapping.values())
            # Explicit metadata validation avoids SQLite interpreting unknown quoted columns as literals.
            known = {r[1] for r in conn.execute('PRAGMA table_info("' + self.table + '")')}
            if not set(self.mapping.values()).issubset(known):
                raise ContractError('MISSING_COLUMN')
            rows = [dict(zip(self.mapping, row)) for row in conn.execute(
                'SELECT ' + columns + ' FROM "' + self.table + '" LIMIT ?', (MAX_ROWS + 1,))]
            if len(rows) > MAX_ROWS:
                raise ContractError('BATCH_TOO_LARGE')
            return SourceBatch(rows, {'kind': 'sqlite', 'locator': str(Path(self.path).resolve()),
                                      'table': self.table, 'mapping': self.mapping,
                                      'sha256': digest(rows), 'digest_scope': 'selected_rows'})
        except ContractError:
            raise
        except (sqlite3.Error, OSError, ValueError, TypeError) as exc:
            raise ContractError('SOURCE_READ_FAILED', type(exc).__name__) from exc
        finally:
            if conn is not None:
                conn.close()


@dataclass(frozen=True)
class Snapshot:
    payload: str

    @property
    def snapshot_id(self):
        return hashlib.sha256(self.payload.encode('utf-8')).hexdigest()

    @property
    def semantic_digest(self):
        doc = self.to_dict()
        return digest({k: doc[k] for k in ('model_ref', 'catalog_digest', 'records')})

    def to_dict(self):
        return json.loads(self.payload)


def load_snapshot(model_path, objects_path, source, expected_model_ref, now):
    model = read_json(model_path)
    if not isinstance(model, dict) or set(model) != {'id', 'version', 'record_schema', 'max_age_seconds'}:
        raise ContractError('MODEL_INVALID')
    ref = {'id': model['id'], 'version': model['version'], 'sha256': digest(model)}
    if ref != expected_model_ref:
        raise ContractError('MODEL_BINDING_MISMATCH')
    if model['version'] != '1.0.0':
        raise ContractError('UNSUPPORTED_MODEL_VERSION')
    age = model['max_age_seconds']
    if type(age) is not int or age <= 0:
        raise ContractError('MODEL_INVALID')
    schema = model['record_schema']
    if not isinstance(schema, dict) or schema.get('type') != 'object' or not isinstance(schema.get('properties'), dict):
        raise ContractError('MODEL_INVALID')
    if not isinstance(source.mapping, dict) or not all(isinstance(v, str) for v in source.mapping.values()):
        raise ContractError('MAPPING_MISMATCH')
    def forbid_refs(value):
        if isinstance(value, dict):
            if any(k in value for k in ('$ref', '$dynamicRef', '$recursiveRef')):
                raise ContractError('SCHEMA_REF_NOT_SUPPORTED')
            for child in value.values():
                forbid_refs(child)
        elif isinstance(value, list):
            for child in value:
                forbid_refs(child)
    forbid_refs(model['record_schema'])
    try:
        Draft202012Validator.check_schema(model['record_schema'])
        validator = Draft202012Validator(model['record_schema'], format_checker=FormatChecker())
    except Exception as exc:
        raise ContractError('MODEL_INVALID') from exc
    objects = read_json(objects_path)
    if not isinstance(objects, list) or not all(isinstance(x, str) for x in objects) or len(objects) != len(set(objects)):
        raise ContractError('OBJECT_CATALOG_INVALID')
    if now.tzinfo is None or now.utcoffset() is None:
        raise ContractError('TIMEZONE_REQUIRED')
    if set(source.mapping) != set(model['record_schema'].get('properties', {})):
        raise ContractError('MAPPING_MISMATCH')
    batch = source.read()
    if not batch.rows:
        raise ContractError('EMPTY_BATCH')
    seen = set()
    normalized = []
    for index, record in enumerate(batch.rows):
        # Do not rely on optional jsonschema date-time extras being installed.
        if not isinstance(record.get('observed_at'), str):
            raise ContractError('RECORD_INVALID', f'row={index + 1}; field=observed_at')
        observed = timestamp(record['observed_at'])
        errors = list(validator.iter_errors(record))
        if errors:
            raise ContractError('RECORD_INVALID', f'row={index + 1}; field={list(errors[0].path)}')
        obj = record.get('object_ref')
        if obj not in objects:
            raise ContractError('UNKNOWN_OBJECT', f'row={index + 1}')
        if obj in seen:
            raise ContractError('DUPLICATE_OBJECT', f'row={index + 1}')
        seen.add(obj)
        elapsed = (now - observed).total_seconds()
        if elapsed < 0:
            raise ContractError('FUTURE_OBSERVATION')
        if elapsed > age:
            raise ContractError('STALE_FACT')
        item = dict(record)
        # Normalise source-specific storage values only under the published schema.
        # This is not an inference or a unit conversion.
        for name, descriptor in schema['properties'].items():
            if descriptor.get('type') == 'integer' and isinstance(item.get(name), float) and item[name].is_integer():
                item[name] = int(item[name])
        item['observed_at'] = observed.isoformat()
        normalized.append(item)
    normalized.sort(key=lambda x: x['object_ref'])
    return Snapshot(canonical({'contract': 'observation-input-spike/1', 'synthetic_demo': True,
                               'model_ref': ref, 'catalog_digest': digest(objects),
                               'records': normalized, 'source': batch.provenance,
                               'validated_at': now.astimezone(timezone.utc).isoformat()}))
